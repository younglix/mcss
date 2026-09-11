from decimal import Decimal

from .models import Discount, Invoice, NonAcademicStaffPayout, Payslip, Scholarship

# Exact column order the bank's own reference template uses, with "Title"
# inserted right after "Names" per spec. Getting this list right is
# everything else below derives its column positions from it — nothing is
# hardcoded to a specific letter, so inserting/reordering a column here
# can't silently point the Narration formula chain (see _build_payout_rows)
# at the wrong cell the way a hardcoded "=H2" would.
PAYOUT_SHEET_HEADERS = [
    "Payment Reference", "Beneficiary Code", "Names", "Title", "Account No.",
    "Account Type", "CBN Sort code", "Is CashCard", "Narration", "Amount",
    "Email Address", "Currency Code",
]
# Matches the reference template's widths 1:1 for every column that exists
# there; "Title" (new) and "Currency Code" (unset in the reference) get a
# reasonable width of their own.
PAYOUT_SHEET_COLUMN_WIDTHS = [20.5, 19.5, 44.5, 20, 16.33, 14.33, 21.16, 15, 18.33, 13.16, 30.33, 12]

_STAFF_CUSTOM_FIELD_KEYS = [
    "payment_reference", "beneficiary_code", "title",
    "account_number", "account_type", "sort_code", "is_cashcard",
]


def apply_scholarship_allocation(allocation, applied_by=None):
    """Discounts every one of the allocation's student's not-yet-settled
    invoices in that session — the actual effect of granting a scholarship.
    Idempotent per invoice (a Discount already linked to this allocation is
    never duplicated), so re-calling this after new invoices appear for the
    same student/session only touches the new ones."""
    scholarship = allocation.scholarship
    invoices = Invoice.objects.filter(
        student=allocation.student, session=allocation.session,
    ).exclude(status__in=[Invoice.Status.PAID, Invoice.Status.WAIVED])

    created = 0
    for invoice in invoices:
        if Discount.objects.filter(invoice=invoice, scholarship_allocation=allocation).exists():
            continue
        if scholarship.coverage_type == Scholarship.CoverageType.PERCENTAGE:
            amount = (invoice.amount * scholarship.coverage_value / Decimal("100")).quantize(Decimal("0.01"))
        else:
            amount = scholarship.coverage_value
        amount = min(amount, invoice.amount)
        Discount.objects.create(
            invoice=invoice, amount=amount, reason=f"Scholarship: {scholarship.name}",
            scholarship_allocation=allocation, applied_by=applied_by,
        )
        invoice.refresh_status()
        created += 1
    return created


def _staff_bank_details():
    """{user_id (str): {custom_field_key: value}} for every active staff
    member — the Title + 6 bank/payout fields, which live as
    apps.custom_fields values (entity="staff"), not a bespoke model, per
    Requirement 1. Keyed by the field's `key` slug so a missing definition
    (Super Admin hasn't added that field yet) just means that column comes
    out blank, not an error."""
    from apps.custom_fields.models import CustomField, CustomFieldValue

    fields_by_id = {
        f.id: f.key
        for f in CustomField.objects.filter(entity=CustomField.Entity.STAFF, key__in=_STAFF_CUSTOM_FIELD_KEYS)
    }
    if not fields_by_id:
        return {}

    by_staff = {}
    for v in CustomFieldValue.objects.filter(field_id__in=fields_by_id).values("field_id", "entity_id", "value"):
        by_staff.setdefault(str(v["entity_id"]), {})[fields_by_id[v["field_id"]]] = v["value"]
    return by_staff


def build_payout_sheet(run, narration):
    """The HR payout sheet — every active regular (system-user) staff
    member's payslip for `run`, plus every active non-academic (payroll-only)
    staff record — written into one .xlsx matching the bank template's exact
    column order, font (Times New Roman 12, bold headers, bold Amount), and
    the Narration "formula chain" (row 2 holds the literal label typed by
    HR; every row after references the cell directly above it, so retyping
    the label once on row 2 updates the whole column).

    `run` may be None — non-academic staff never depend on a PayrollRun at
    all, and this still produces a complete sheet for them alone (regular
    staff rows all get a blank Amount in that case, same as "no payslip").

    Returns (workbook, stats): stats reports how many rows came from each
    source and how many regular-staff rows had no payslip for this run
    (Amount left blank for those — they're still included, never crash).
    """
    import openpyxl
    from django.contrib.auth import get_user_model
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    User = get_user_model()

    narration_col = PAYOUT_SHEET_HEADERS.index("Narration") + 1
    amount_col = PAYOUT_SHEET_HEADERS.index("Amount") + 1
    narration_letter = get_column_letter(narration_col)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Payout"

    header_font = Font(name="Times New Roman", size=12, bold=True)
    data_font = Font(name="Times New Roman", size=12, bold=False)
    amount_font = Font(name="Times New Roman", size=12, bold=True)

    for col, header in enumerate(PAYOUT_SHEET_HEADERS, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_font
    for col, width in enumerate(PAYOUT_SHEET_COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = width

    def write_row(row, *, payment_reference, beneficiary_code, names, title, account_no,
                  account_type, sort_code, is_cashcard, amount, email, currency_code):
        values = {
            "Payment Reference": payment_reference, "Beneficiary Code": beneficiary_code,
            "Names": names, "Title": title, "Account No.": account_no,
            "Account Type": account_type, "CBN Sort code": sort_code,
            "Is CashCard": int(bool(is_cashcard)),
            "Narration": narration if row == 2 else f"={narration_letter}{row - 1}",
            "Amount": amount, "Email Address": email, "Currency Code": currency_code,
        }
        for col, header in enumerate(PAYOUT_SHEET_HEADERS, start=1):
            cell = ws.cell(row=row, column=col, value=values[header])
            cell.font = amount_font if col == amount_col else data_font
        ws.cell(row=row, column=amount_col).number_format = "#,##0.00"

    bank_details = _staff_bank_details()
    payslips_by_staff = {}
    if run is not None:
        payslips_by_staff = {
            str(p.staff_id): p.net_pay for p in Payslip.objects.filter(run=run).select_related("staff")
        }

    row = 2
    regular_count = 0
    missing_payslip_count = 0
    staff_qs = User.objects.filter(user_type=User.UserType.STAFF, is_active=True, is_deleted=False).order_by("full_name")
    for user in staff_qs:
        details = bank_details.get(str(user.id), {})
        amount = payslips_by_staff.get(str(user.id))
        if amount is None:
            missing_payslip_count += 1
        write_row(
            row, payment_reference=details.get("payment_reference") or "",
            beneficiary_code=details.get("beneficiary_code") or "", names=user.full_name,
            title=details.get("title") or "", account_no=details.get("account_number") or "",
            account_type=details.get("account_type") or "", sort_code=details.get("sort_code") or "",
            is_cashcard=details.get("is_cashcard"), amount=amount, email=user.email or "", currency_code="NGN",
        )
        row += 1
        regular_count += 1

    non_academic_count = 0
    for rec in NonAcademicStaffPayout.objects.filter(is_active=True).order_by("full_name"):
        write_row(
            row, payment_reference=rec.payment_reference, beneficiary_code=rec.beneficiary_code,
            names=rec.full_name, title=rec.title, account_no=rec.account_number,
            account_type=rec.account_type, sort_code=rec.sort_code, is_cashcard=rec.is_cashcard,
            amount=rec.pay_amount or None, email=rec.email, currency_code=rec.currency_code,
        )
        row += 1
        non_academic_count += 1

    stats = {
        "regular_staff": regular_count,
        "non_academic_staff": non_academic_count,
        "total_rows": regular_count + non_academic_count,
        "missing_payslip": missing_payslip_count,
    }
    return wb, stats
