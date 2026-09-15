"""HR Payout Sheet: exact column order (Title right after Names), the
Narration formula chain, Amount sourced from Payslip (regular staff) vs a
plain pay figure (non-academic staff), active-only filtering, and the
export permission gate.
"""

import hashlib
import hmac
import json
from decimal import Decimal
from unittest.mock import patch

from django.contrib.auth import get_user_model
from django.test import TestCase
from openpyxl.utils import get_column_letter
from rest_framework.test import APIClient

from apps.custom_fields.models import CustomField, CustomFieldValue
from apps.rbac.models import Permission, Role, RolePermission, UserRole
from apps.settings_app.models import SystemSetting

from . import services
from .models import NonAcademicStaffPayout, PayrollRun, Payslip
from .services import PAYOUT_SHEET_HEADERS

User = get_user_model()


class PayoutSheetTestBase(TestCase):
    def setUp(self):
        self.title_field = CustomField.objects.create(entity=CustomField.Entity.STAFF, key="title", label="Title")
        self.acct_field = CustomField.objects.create(entity=CustomField.Entity.STAFF, key="account_number", label="Account No.")
        self.run = PayrollRun.objects.create(month=1, year=2199)  # arbitrary, unused elsewhere

    def make_staff(self, name, *, active=True, title=None, account_number=None, net_pay=None):
        user = User.objects.create(full_name=name, email=f"{name.lower().replace(' ', '.')}@x.io", user_type="staff", is_active=active)
        if title:
            CustomFieldValue.objects.create(field=self.title_field, entity_id=user.id, value=title)
        if account_number:
            CustomFieldValue.objects.create(field=self.acct_field, entity_id=user.id, value=account_number)
        if net_pay is not None:
            Payslip.objects.create(
                run=self.run, staff=user, basic_salary=net_pay, allowances=Decimal("0"),
                deductions=Decimal("0"), net_pay=net_pay,
            )
        return user


class ColumnOrderTests(PayoutSheetTestBase):
    def test_header_order_matches_spec_exactly(self):
        wb, _ = services.build_payout_sheet(self.run, "TEST SALARY")
        ws = wb.active
        headers = [ws.cell(row=1, column=c).value for c in range(1, len(PAYOUT_SHEET_HEADERS) + 1)]
        self.assertEqual(headers, [
            "Payment Reference", "Beneficiary Code", "Names", "Title", "Account No.",
            "Account Type", "CBN Sort code", "Is CashCard", "Narration", "Amount",
            "Email Address", "Currency Code",
        ])

    def test_title_sits_right_after_names(self):
        names_col = PAYOUT_SHEET_HEADERS.index("Names")
        title_col = PAYOUT_SHEET_HEADERS.index("Title")
        self.assertEqual(title_col, names_col + 1)

    def test_header_and_amount_styling(self):
        self.make_staff("Font Check", net_pay=Decimal("50000"))
        wb, _ = services.build_payout_sheet(self.run, "TEST SALARY")
        ws = wb.active
        header_font = ws.cell(row=1, column=1).font
        self.assertEqual(header_font.name, "Times New Roman")
        self.assertEqual(header_font.size, 12)
        self.assertTrue(header_font.bold)

        amount_col = PAYOUT_SHEET_HEADERS.index("Amount") + 1
        names_col = PAYOUT_SHEET_HEADERS.index("Names") + 1
        self.assertTrue(ws.cell(row=2, column=amount_col).font.bold)
        self.assertEqual(ws.cell(row=2, column=amount_col).number_format, "#,##0.00")
        self.assertFalse(ws.cell(row=2, column=names_col).font.bold)


class NarrationChainTests(PayoutSheetTestBase):
    def test_first_row_is_literal_rest_reference_the_row_above(self):
        self.make_staff("Person A", net_pay=Decimal("10000"))
        self.make_staff("Person B", net_pay=Decimal("20000"))
        self.make_staff("Person C", net_pay=Decimal("30000"))
        wb, stats = services.build_payout_sheet(self.run, "SEPTEMBER SALARY")
        ws = wb.active
        narration_col = PAYOUT_SHEET_HEADERS.index("Narration") + 1
        narration_letter = get_column_letter(narration_col)

        self.assertEqual(ws.cell(row=2, column=narration_col).value, "SEPTEMBER SALARY")
        for r in range(3, 2 + stats["total_rows"]):
            self.assertEqual(ws.cell(row=r, column=narration_col).value, f"={narration_letter}{r - 1}")

    def test_chain_points_at_narration_column_not_a_hardcoded_letter(self):
        """Regression guard: Title shifts every column after it one place to
        the right versus the original 11-column bank template, where the
        formula chain was literally "=H2". If this ever hardcoded "H" again
        it would silently point at Is CashCard instead of Narration."""
        self.make_staff("Person A", net_pay=Decimal("10000"))
        self.make_staff("Person B", net_pay=Decimal("20000"))
        wb, _ = services.build_payout_sheet(self.run, "X")
        ws = wb.active
        formula = ws.cell(row=3, column=PAYOUT_SHEET_HEADERS.index("Narration") + 1).value
        self.assertTrue(formula.startswith("=I"), f"expected a reference into column I (Narration), got {formula!r}")


class AmountSourcingTests(PayoutSheetTestBase):
    def test_amount_from_payslip_when_present(self):
        self.make_staff("Has Payslip", net_pay=Decimal("115000.00"))
        wb, _ = services.build_payout_sheet(self.run, "X")
        ws = wb.active
        amount_col = PAYOUT_SHEET_HEADERS.index("Amount") + 1
        self.assertEqual(ws.cell(row=2, column=amount_col).value, Decimal("115000.00"))

    def test_amount_blank_and_included_when_no_payslip(self):
        self.make_staff("No Payslip At All")
        wb, stats = services.build_payout_sheet(self.run, "X")
        ws = wb.active
        amount_col = PAYOUT_SHEET_HEADERS.index("Amount") + 1
        names_col = PAYOUT_SHEET_HEADERS.index("Names") + 1
        self.assertEqual(ws.cell(row=2, column=names_col).value, "No Payslip At All")
        self.assertIsNone(ws.cell(row=2, column=amount_col).value)
        self.assertEqual(stats["missing_payslip"], 1)

    def test_non_academic_staff_amount_is_the_plain_pay_figure_no_run_needed(self):
        NonAcademicStaffPayout.objects.create(full_name="A Cleaner", title="Cleaner", pay_amount="35000", is_active=True)
        wb, stats = services.build_payout_sheet(None, "X")  # no run at all
        ws = wb.active
        amount_col = PAYOUT_SHEET_HEADERS.index("Amount") + 1
        self.assertEqual(ws.cell(row=2, column=amount_col).value, "35000")
        self.assertEqual(stats["non_academic_staff"], 1)
        self.assertEqual(stats["missing_payslip"], 0)  # non-academic rows never count against this


class ActiveOnlyFilterTests(PayoutSheetTestBase):
    def test_inactive_staff_excluded(self):
        self.make_staff("Active One", active=True, net_pay=Decimal("1000"))
        self.make_staff("Inactive One", active=False, net_pay=Decimal("1000"))
        wb, stats = services.build_payout_sheet(self.run, "X")
        ws = wb.active
        names_col = PAYOUT_SHEET_HEADERS.index("Names") + 1
        names = [ws.cell(row=r, column=names_col).value for r in range(2, 2 + stats["total_rows"])]
        self.assertIn("Active One", names)
        self.assertNotIn("Inactive One", names)

    def test_inactive_non_academic_staff_excluded(self):
        NonAcademicStaffPayout.objects.create(full_name="Left Already", pay_amount="0", is_active=False)
        wb, stats = services.build_payout_sheet(None, "X")
        self.assertEqual(stats["non_academic_staff"], 0)

    def test_new_staff_appears_with_no_code_change(self):
        """"As new staff are recruited, they appear automatically" — a
        freshly created active staff user, never explicitly wired into
        anything, shows up on the very next generation."""
        self.make_staff("Brand New Hire", active=True)
        wb, stats = services.build_payout_sheet(self.run, "X")
        ws = wb.active
        names_col = PAYOUT_SHEET_HEADERS.index("Names") + 1
        names = [ws.cell(row=r, column=names_col).value for r in range(2, 2 + stats["total_rows"])]
        self.assertIn("Brand New Hire", names)


class TitleAndBankFieldSourcingTests(PayoutSheetTestBase):
    def test_title_and_bank_fields_come_from_custom_field_values(self):
        self.make_staff("Full Details", title="Bursar", account_number="0123456789", net_pay=Decimal("1000"))
        wb, _ = services.build_payout_sheet(self.run, "X")
        ws = wb.active
        title_col = PAYOUT_SHEET_HEADERS.index("Title") + 1
        acct_col = PAYOUT_SHEET_HEADERS.index("Account No.") + 1
        self.assertEqual(ws.cell(row=2, column=title_col).value, "Bursar")
        self.assertEqual(ws.cell(row=2, column=acct_col).value, "0123456789")

    def test_missing_custom_field_values_render_blank_not_crash(self):
        self.make_staff("No Bank Details Yet", net_pay=Decimal("1000"))
        wb, _ = services.build_payout_sheet(self.run, "X")  # must not raise
        ws = wb.active
        title_col = PAYOUT_SHEET_HEADERS.index("Title") + 1
        self.assertEqual(ws.cell(row=2, column=title_col).value, "")


class PermissionGateTests(PayoutSheetTestBase):
    def setUp(self):
        super().setUp()
        role = Role.objects.create(name="HR Export", slug="hr-export-test", is_system=False)
        perm, _ = Permission.objects.get_or_create(code="payroll.export", defaults={"module": "payroll", "action": "export"})
        RolePermission.objects.create(role=role, permission=perm)
        self.hr_user = User.objects.create(full_name="HR Exporter", email="hrexp@x.io", user_type="staff", is_active=True)
        UserRole.objects.create(user=self.hr_user, role=role)

        self.no_perm_user = User.objects.create(full_name="No Perm", email="noperm2@x.io", user_type="staff", is_active=True)
        self.client = APIClient()

    def test_hr_with_export_permission_can_download(self):
        self.client.force_authenticate(self.hr_user)
        res = self.client.get(f"/api/v1/finance/payroll/payout-sheet?run={self.run.id}&narration=TEST")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res["Content-Type"], "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("X-Payout-Total-Rows", res)

    def test_staff_without_the_permission_is_blocked(self):
        self.client.force_authenticate(self.no_perm_user)
        res = self.client.get(f"/api/v1/finance/payroll/payout-sheet?run={self.run.id}&narration=TEST")
        self.assertEqual(res.status_code, 403)

    def test_narration_is_required(self):
        self.client.force_authenticate(self.hr_user)
        res = self.client.get(f"/api/v1/finance/payroll/payout-sheet?run={self.run.id}")
        self.assertEqual(res.status_code, 400)


class FeeItemSelfServiceTests(TestCase):
    """A student browsing Fee Items and self-creating a ticket (Invoice) for
    an optional one-off fee — the flow InvoicePayView's own docstring already
    anticipated ("later optional/add-on fees"), now wired up end to end."""

    def setUp(self):
        from datetime import date

        from apps.academics.models import Student
        from apps.configuration.models import AcademicSession, FeeCategory

        self.session = AcademicSession.objects.create(
            name="2026/2027", start_date=date(2026, 9, 1), end_date=date(2027, 7, 31), is_current=True,
        )
        self.sportswear = FeeCategory.objects.create(name="Sportswear", amount=Decimal("15000"), is_recurring=False)
        self.tuition = FeeCategory.objects.create(name="Tuition", amount=Decimal("100000"), is_recurring=True)

        self.student_user = User.objects.create(full_name="Student One", email="student1@x.io", user_type="student", is_active=True)
        self.student = Student.objects.create(user=self.student_user, date_of_birth=date(2012, 1, 1))
        self.other_staff = User.objects.create(full_name="Plain Staff", email="staff1@x.io", user_type="staff", is_active=True)
        self.client = APIClient()

    def test_student_sees_only_non_recurring_items(self):
        self.client.force_authenticate(self.student_user)
        res = self.client.get("/api/v1/finance/fee-items/mine")
        self.assertEqual(res.status_code, 200)
        names = [f["name"] for f in res.json()["data"]]
        self.assertIn("Sportswear", names)
        self.assertNotIn("Tuition", names)

    def test_non_student_cannot_see_the_catalog(self):
        self.client.force_authenticate(self.other_staff)
        res = self.client.get("/api/v1/finance/fee-items/mine")
        self.assertEqual(res.status_code, 403)

    def test_anonymous_is_rejected(self):
        res = self.client.get("/api/v1/finance/fee-items/mine")
        self.assertEqual(res.status_code, 401)

    def test_purchase_creates_a_real_invoice_owned_by_the_student(self):
        self.client.force_authenticate(self.student_user)
        res = self.client.post(f"/api/v1/finance/fee-items/{self.sportswear.id}/purchase", {}, format="json")
        self.assertEqual(res.status_code, 201, res.json())
        body = res.json()["data"]
        self.assertEqual(body["description"], "Sportswear")
        self.assertEqual(Decimal(body["amount"]), Decimal("15000"))
        self.assertEqual(body["status"], "unpaid")

        from .models import Invoice
        invoice = Invoice.objects.get(id=body["id"])
        self.assertEqual(invoice.student_id, self.student.id)
        self.assertEqual(invoice.session_id, self.session.id)
        self.assertEqual(invoice.category_id, self.sportswear.id)

    def test_recurring_item_cannot_be_self_purchased(self):
        self.client.force_authenticate(self.student_user)
        res = self.client.post(f"/api/v1/finance/fee-items/{self.tuition.id}/purchase", {}, format="json")
        self.assertEqual(res.status_code, 400)

    def test_clicking_twice_does_not_create_a_duplicate_open_ticket(self):
        self.client.force_authenticate(self.student_user)
        first = self.client.post(f"/api/v1/finance/fee-items/{self.sportswear.id}/purchase", {}, format="json")
        second = self.client.post(f"/api/v1/finance/fee-items/{self.sportswear.id}/purchase", {}, format="json")
        self.assertEqual(first.json()["data"]["id"], second.json()["data"]["id"])
        self.assertTrue(first.json()["data"]["created"])
        self.assertFalse(second.json()["data"]["created"])

        from .models import Invoice
        self.assertEqual(Invoice.objects.filter(student=self.student, description="Sportswear").count(), 1)

    def test_a_new_ticket_can_be_created_after_the_first_is_paid(self):
        from .models import Invoice, Payment

        self.client.force_authenticate(self.student_user)
        first = self.client.post(f"/api/v1/finance/fee-items/{self.sportswear.id}/purchase", {}, format="json")
        invoice = Invoice.objects.get(id=first.json()["data"]["id"])
        Payment.objects.create(invoice=invoice, amount=Decimal("15000"))
        invoice.refresh_status()
        self.assertEqual(invoice.status, Invoice.Status.PAID)

        second = self.client.post(f"/api/v1/finance/fee-items/{self.sportswear.id}/purchase", {}, format="json")
        self.assertEqual(second.status_code, 201)
        self.assertNotEqual(second.json()["data"]["id"], first.json()["data"]["id"])

    def test_purchased_ticket_appears_in_the_students_own_invoice_list(self):
        self.client.force_authenticate(self.student_user)
        self.client.post(f"/api/v1/finance/fee-items/{self.sportswear.id}/purchase", {}, format="json")
        res = self.client.get("/api/v1/finance/invoices/mine")
        descriptions = [i["description"] for i in res.json()["data"]]
        self.assertIn("Sportswear", descriptions)

    def test_non_student_cannot_purchase(self):
        self.client.force_authenticate(self.other_staff)
        res = self.client.post(f"/api/v1/finance/fee-items/{self.sportswear.id}/purchase", {}, format="json")
        self.assertEqual(res.status_code, 403)


class RestrictionCheckServiceTests(TestCase):
    """apps.finance.services.restriction_check is the one shared definition
    of "is this student blocked from X" every gated endpoint (library,
    hostel, transport, activities, report-card downloads, exam attempts)
    calls — covered here in isolation from any of those call sites."""

    def setUp(self):
        from datetime import date

        from apps.academics.models import Student
        from apps.configuration.models import AcademicSession, FeeCategory

        self.session = AcademicSession.objects.create(
            name="2026/2027", start_date=date(2026, 9, 1), end_date=date(2027, 7, 31), is_current=True,
        )
        self.other_session = AcademicSession.objects.create(
            name="2025/2026", start_date=date(2025, 9, 1), end_date=date(2026, 7, 31), is_current=False,
        )
        self.library_fee = FeeCategory.objects.create(name="Library Fee", is_recurring=True, restriction_type=FeeCategory.RestrictionType.LIBRARY)
        self.plain_fee = FeeCategory.objects.create(name="Textbooks", is_recurring=False, amount=Decimal("5000"))
        student_user = User.objects.create(full_name="Student One", email="stu1@x.io", user_type="student", is_active=True)
        self.student = Student.objects.create(user=student_user)

    def _invoice(self, category, session=None, amount=Decimal("3000")):
        from .models import Invoice

        return Invoice.objects.create(
            student=self.student, session=session or self.session, category=category,
            description=category.name, amount=amount,
        )

    def test_no_outstanding_invoice_means_not_blocked(self):
        from apps.configuration.models import FeeCategory
        from . import services

        blocked, invoices, message = services.restriction_check(self.student, FeeCategory.RestrictionType.LIBRARY)
        self.assertFalse(blocked)
        self.assertEqual(invoices, [])
        self.assertEqual(message, "")

    def test_unpaid_invoice_against_a_restricted_category_blocks(self):
        from apps.configuration.models import FeeCategory
        from . import services

        self._invoice(self.library_fee)
        blocked, invoices, message = services.restriction_check(self.student, FeeCategory.RestrictionType.LIBRARY)
        self.assertTrue(blocked)
        self.assertEqual(len(invoices), 1)
        self.assertIn("Library Fee", message)

    def test_paying_it_off_clears_the_restriction(self):
        from apps.configuration.models import FeeCategory
        from . import services
        from .models import Payment

        invoice = self._invoice(self.library_fee)
        Payment.objects.create(invoice=invoice, amount=invoice.amount)
        invoice.refresh_status()

        blocked, _invoices, _message = services.restriction_check(self.student, FeeCategory.RestrictionType.LIBRARY)
        self.assertFalse(blocked)

    def test_a_waived_invoice_does_not_block(self):
        from apps.configuration.models import FeeCategory
        from . import services
        from .models import Invoice

        invoice = self._invoice(self.library_fee)
        invoice.status = Invoice.Status.WAIVED
        invoice.save(update_fields=["status"])

        blocked, _invoices, _message = services.restriction_check(self.student, FeeCategory.RestrictionType.LIBRARY)
        self.assertFalse(blocked)

    def test_a_category_with_no_restriction_type_never_blocks_anything(self):
        from apps.configuration.models import FeeCategory
        from . import services

        self._invoice(self.plain_fee)
        blocked, _invoices, _message = services.restriction_check(self.student, FeeCategory.RestrictionType.LIBRARY)
        self.assertFalse(blocked)

    def test_an_outstanding_invoice_of_a_different_restriction_type_does_not_block(self):
        from apps.configuration.models import FeeCategory
        from . import services

        hostel_fee = FeeCategory.objects.create(name="Hostel Accommodation Fee", is_recurring=True, restriction_type=FeeCategory.RestrictionType.HOSTEL)
        self._invoice(hostel_fee)
        blocked, _invoices, _message = services.restriction_check(self.student, FeeCategory.RestrictionType.LIBRARY)
        self.assertFalse(blocked)

    def test_a_prior_sessions_unpaid_invoice_does_not_reach_forward(self):
        from apps.configuration.models import FeeCategory
        from . import services

        self._invoice(self.library_fee, session=self.other_session)
        blocked, _invoices, _message = services.restriction_check(self.student, FeeCategory.RestrictionType.LIBRARY)
        self.assertFalse(blocked)


class RestrictionSummaryViewTests(TestCase):
    """MyRestrictionsView / ChildRestrictionsView — the Fees & Receipts
    banner data for the student and parent portals."""

    def setUp(self):
        from datetime import date

        from apps.academics.models import Student
        from apps.configuration.models import AcademicSession, FeeCategory

        self.session = AcademicSession.objects.create(
            name="2026/2027", start_date=date(2026, 9, 1), end_date=date(2027, 7, 31), is_current=True,
        )
        self.library_fee = FeeCategory.objects.create(name="Library Fee", is_recurring=True, restriction_type=FeeCategory.RestrictionType.LIBRARY)
        student_user = User.objects.create(full_name="Student One", email="stu1@x.io", user_type="student", is_active=True)
        guardian_user = User.objects.create(full_name="Guardian One", email="guardian1@x.io", user_type="parent", is_active=True)
        self.student = Student.objects.create(user=student_user, guardian_user=guardian_user)
        self.guardian_user = guardian_user
        self.client = APIClient()

    def test_no_restrictions_when_nothing_is_outstanding(self):
        self.client.force_authenticate(self.student.user)
        res = self.client.get("/api/v1/finance/restrictions/mine")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json()["data"], [])

    def test_lists_an_active_restriction_with_its_reason(self):
        from .models import Invoice

        Invoice.objects.create(
            student=self.student, session=self.session, category=self.library_fee,
            description="Library Fee", amount=Decimal("3000"),
        )
        self.client.force_authenticate(self.student.user)
        res = self.client.get("/api/v1/finance/restrictions/mine")
        data = res.json()["data"]
        self.assertEqual(len(data), 1)
        self.assertEqual(data[0]["restriction_type"], "library")
        self.assertIn("Library Fee", data[0]["message"])

    def test_guardian_sees_the_same_restriction_for_their_child(self):
        from .models import Invoice

        Invoice.objects.create(
            student=self.student, session=self.session, category=self.library_fee,
            description="Library Fee", amount=Decimal("3000"),
        )
        self.client.force_authenticate(self.guardian_user)
        res = self.client.get(f"/api/v1/finance/restrictions/child/{self.student.id}")
        self.assertEqual(len(res.json()["data"]), 1)

    def test_a_stranger_cannot_see_another_students_restrictions(self):
        outsider = User.objects.create(full_name="Outsider", email="outsider@x.io", user_type="parent", is_active=True)
        self.client.force_authenticate(outsider)
        res = self.client.get(f"/api/v1/finance/restrictions/child/{self.student.id}")
        self.assertEqual(res.status_code, 403)


class PaystackPaymentTestBase(TestCase):
    """Covers the actual bug report: a real successful Paystack payment
    left the ticket unpaid and no Payment record behind, because nothing
    ever verified the transaction server-side — the webhook needs a URL
    registered in the Paystack dashboard and reachable from Paystack's
    servers (neither guaranteed, especially with no domain/HTTPS yet), and
    nothing else called verify_transaction() at all. Fixed with a shared
    services.record_paystack_payment(), called from both the webhook
    (PaystackWebhookView) and a new client-triggered endpoint
    (PaystackVerifyReturnView) the student's browser hits the moment it
    lands back on /student/finance."""

    def setUp(self):
        from datetime import date

        from apps.academics.models import Student
        from apps.configuration.models import AcademicSession

        self.session = AcademicSession.objects.create(
            name="2026/2027", start_date=date(2026, 9, 1), end_date=date(2027, 7, 31), is_current=True,
        )
        secret_setting = SystemSetting(key="payments.paystack.secret_key", group="payments", is_secret=True)
        secret_setting.set_value("sk_test_fakefakefake")
        secret_setting.save()

        self.student_user = User.objects.create(full_name="Student One", email="student1@x.io", user_type="student", is_active=True)
        self.student = Student.objects.create(user=self.student_user, date_of_birth=date(2012, 1, 1))
        self.other_student_user = User.objects.create(full_name="Student Two", email="student2@x.io", user_type="student", is_active=True)
        self.other_student = Student.objects.create(user=self.other_student_user, date_of_birth=date(2012, 1, 1))
        self.staff_user = User.objects.create(full_name="Staff One", email="staff1@x.io", user_type="staff", is_active=True)

        from .models import Invoice

        self.invoice = Invoice.objects.create(
            student=self.student, description="Sportswear", session=self.session, amount=Decimal("15000"),
        )
        self.reference = f"MCSS-{self.invoice.id.hex}-abcd1234"
        self.client = APIClient()

    def paystack_success_payload(self, amount_naira="15000"):
        return {"status": "success", "amount": int(Decimal(amount_naira) * 100), "reference": self.reference}


class RecordPaystackPaymentServiceTests(PaystackPaymentTestBase):
    @patch("apps.admissions.paystack.verify_transaction")
    def test_creates_a_completed_payment_and_marks_the_invoice_paid(self, mock_verify):
        mock_verify.return_value = self.paystack_success_payload()

        from .models import Invoice, Payment

        payment, created, error = services.record_paystack_payment(self.reference)
        self.assertIsNone(error)
        self.assertTrue(created)
        self.assertEqual(payment.status, Payment.Status.COMPLETED)
        self.assertEqual(payment.method, Payment.Method.ONLINE)
        self.assertEqual(payment.amount, Decimal("15000"))

        self.invoice.refresh_from_db()
        self.assertEqual(self.invoice.status, Invoice.Status.PAID)

    @patch("apps.admissions.paystack.verify_transaction")
    def test_is_idempotent_across_repeated_calls(self, mock_verify):
        mock_verify.return_value = self.paystack_success_payload()

        from .models import Payment

        first_payment, first_created, _ = services.record_paystack_payment(self.reference)
        second_payment, second_created, _ = services.record_paystack_payment(self.reference)

        self.assertTrue(first_created)
        self.assertFalse(second_created)
        self.assertEqual(first_payment.id, second_payment.id)
        self.assertEqual(Payment.objects.filter(reference=self.reference).count(), 1)

    @patch("apps.admissions.paystack.verify_transaction")
    def test_rejects_a_transaction_paystack_does_not_confirm_as_successful(self, mock_verify):
        mock_verify.return_value = {"status": "failed", "amount": 1500000}

        from .models import Payment

        payment, created, error = services.record_paystack_payment(self.reference)
        self.assertIsNone(payment)
        self.assertFalse(created)
        self.assertIsNotNone(error)
        self.assertFalse(Payment.objects.filter(reference=self.reference).exists())

    @patch("apps.admissions.paystack.verify_transaction")
    def test_rejects_when_paystack_verification_itself_fails(self, mock_verify):
        mock_verify.return_value = None  # network error / bad reference, see paystack.verify_transaction
        payment, created, error = services.record_paystack_payment(self.reference)
        self.assertIsNone(payment)
        self.assertFalse(created)
        self.assertIsNotNone(error)

    def test_rejects_a_missing_reference(self):
        payment, created, error = services.record_paystack_payment("")
        self.assertIsNone(payment)
        self.assertFalse(created)
        self.assertIsNotNone(error)

    @patch("apps.admissions.paystack.verify_transaction")
    def test_the_webhook_path_and_the_client_verify_path_converge_on_one_payment(self, mock_verify):
        """The exact scenario the bug report describes, proven from both
        directions: whichever path (webhook or the browser's own
        return-from-checkout call) reaches the server first, the other is
        a safe no-op — never a duplicate Payment, never a missed one."""
        mock_verify.return_value = self.paystack_success_payload()

        from .models import Invoice, Payment

        # Simulate the client-triggered path winning the race.
        services.record_paystack_payment(self.reference)
        # The webhook redelivers moments later — same underlying call.
        payment, created, error = services.record_paystack_payment(self.reference)

        self.assertIsNone(error)
        self.assertFalse(created)
        self.assertEqual(Payment.objects.filter(reference=self.reference).count(), 1)
        self.invoice.refresh_from_db()
        self.assertEqual(self.invoice.status, Invoice.Status.PAID)


class PaystackWebhookViewTests(PaystackPaymentTestBase):
    def post_webhook(self, payload):
        raw_body = json.dumps(payload).encode("utf-8")
        signature = hmac.new(b"sk_test_fakefakefake", raw_body, hashlib.sha512).hexdigest()
        return self.client.post(
            "/api/v1/finance/payments/paystack/webhook", data=raw_body,
            content_type="application/json", HTTP_X_PAYSTACK_SIGNATURE=signature,
        )

    @patch("apps.admissions.paystack.verify_transaction")
    def test_valid_webhook_creates_the_payment_and_updates_the_invoice(self, mock_verify):
        mock_verify.return_value = self.paystack_success_payload()

        from .models import Invoice, Payment

        res = self.post_webhook({"event": "charge.success", "data": {"reference": self.reference}})
        self.assertEqual(res.status_code, 200, res.json())
        self.assertTrue(Payment.objects.filter(reference=self.reference).exists())
        self.invoice.refresh_from_db()
        self.assertEqual(self.invoice.status, Invoice.Status.PAID)

    def test_an_invalid_signature_is_rejected_without_touching_anything(self):
        from .models import Payment

        raw_body = json.dumps({"event": "charge.success", "data": {"reference": self.reference}}).encode("utf-8")
        res = self.client.post(
            "/api/v1/finance/payments/paystack/webhook", data=raw_body,
            content_type="application/json", HTTP_X_PAYSTACK_SIGNATURE="not-the-real-signature",
        )
        self.assertEqual(res.status_code, 401)
        self.assertFalse(Payment.objects.filter(reference=self.reference).exists())

    def test_a_non_charge_success_event_is_ignored(self):
        res = self.post_webhook({"event": "transfer.success", "data": {"reference": self.reference}})
        self.assertEqual(res.status_code, 200)

        from .models import Payment
        self.assertFalse(Payment.objects.filter(reference=self.reference).exists())

    @patch("apps.admissions.paystack.verify_transaction")
    def test_a_redelivered_webhook_does_not_double_record(self, mock_verify):
        mock_verify.return_value = self.paystack_success_payload()

        from .models import Payment

        self.post_webhook({"event": "charge.success", "data": {"reference": self.reference}})
        second = self.post_webhook({"event": "charge.success", "data": {"reference": self.reference}})
        self.assertEqual(second.status_code, 200)
        self.assertEqual(Payment.objects.filter(reference=self.reference).count(), 1)


class PaystackVerifyReturnViewTests(PaystackPaymentTestBase):
    @patch("apps.admissions.paystack.verify_transaction")
    def test_student_verifying_their_own_reference_marks_the_ticket_paid(self, mock_verify):
        mock_verify.return_value = self.paystack_success_payload()

        from .models import Invoice, Payment

        self.client.force_authenticate(self.student_user)
        res = self.client.post("/api/v1/finance/payments/paystack/verify", {"reference": self.reference}, format="json")
        self.assertEqual(res.status_code, 200, res.json())
        self.assertEqual(res.json()["data"]["status"], "paid")
        self.assertTrue(Payment.objects.filter(reference=self.reference).exists())

    @patch("apps.admissions.paystack.verify_transaction")
    def test_calling_it_twice_is_a_harmless_no_op_the_second_time(self, mock_verify):
        mock_verify.return_value = self.paystack_success_payload()

        from .models import Payment

        self.client.force_authenticate(self.student_user)
        self.client.post("/api/v1/finance/payments/paystack/verify", {"reference": self.reference}, format="json")
        second = self.client.post("/api/v1/finance/payments/paystack/verify", {"reference": self.reference}, format="json")
        self.assertEqual(second.status_code, 200)
        self.assertEqual(Payment.objects.filter(reference=self.reference).count(), 1)

    @patch("apps.admissions.paystack.verify_transaction")
    def test_a_different_students_reference_is_refused_but_the_payment_still_records(self, mock_verify):
        """The payment is real money that actually moved — it still has to
        be recorded against the correct invoice — but the response to the
        wrong caller must not hand back someone else's invoice data."""
        mock_verify.return_value = self.paystack_success_payload()

        from .models import Invoice, Payment

        self.client.force_authenticate(self.other_student_user)
        res = self.client.post("/api/v1/finance/payments/paystack/verify", {"reference": self.reference}, format="json")
        self.assertEqual(res.status_code, 403)
        self.assertTrue(Payment.objects.filter(reference=self.reference).exists())
        self.invoice.refresh_from_db()
        self.assertEqual(self.invoice.status, Invoice.Status.PAID)

    def test_requires_authentication(self):
        res = self.client.post("/api/v1/finance/payments/paystack/verify", {"reference": self.reference}, format="json")
        self.assertEqual(res.status_code, 401)

    def test_a_staff_member_is_rejected_not_just_a_wrong_student(self):
        self.client.force_authenticate(self.staff_user)
        res = self.client.post("/api/v1/finance/payments/paystack/verify", {"reference": self.reference}, format="json")
        self.assertEqual(res.status_code, 403)

    def test_a_missing_reference_is_a_clean_400_not_a_500(self):
        self.client.force_authenticate(self.student_user)
        res = self.client.post("/api/v1/finance/payments/paystack/verify", {}, format="json")
        self.assertEqual(res.status_code, 400)
